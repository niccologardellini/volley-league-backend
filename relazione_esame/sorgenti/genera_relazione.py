"""Genera Relazione_VolleyLeague.pptx (10 slide) e, se LibreOffice e' disponibile, il PDF.

Uso:  python3 diagrammi.py && python3 genera_relazione.py
"""
import datetime
import os
import shutil
import subprocess

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

QUI = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(QUI, "img")
OUT = os.path.abspath(os.path.join(QUI, ".."))
NOME = "Relazione_VolleyLeague"

# Da personalizzare prima della consegna
STUDENTE = "Nome Cognome"
MATRICOLA = "Matricola 000000"
CORSO = "Ingegneria del Software · UniFE · A.A. 2025/26"

NAVY = RGBColor(0x12, 0x1C, 0x46)
BLU = RGBColor(0x1F, 0x4F, 0xD1)
ARANCIO = RGBColor(0xFF, 0x9D, 0x2E)
TESTO = RGBColor(0x1A, 0x20, 0x33)
GRIGIO = RGBColor(0x5A, 0x63, 0x78)
CHIARO = RGBColor(0xF3, 0xF5, 0xFA)
BIANCO = RGBColor(0xFF, 0xFF, 0xFF)
VERDE = RGBColor(0x11, 0x63, 0x36)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
VUOTO = prs.slide_layouts[6]


def img(nome):
    return os.path.join(IMG, nome)


def rett(slide, x, y, w, h, colore, forma=MSO_SHAPE.RECTANGLE, bordo=None):
    s = slide.shapes.add_shape(forma, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = colore
    if bordo:
        s.line.color.rgb = bordo
        s.line.width = Pt(1)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if forma == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    return s


def testo(slide, x, y, w, h, righe, size=14, colore=TESTO, bold=False, align=PP_ALIGN.LEFT,
          anchor=MSO_ANCHOR.TOP, spazio=4):
    """righe: stringa o lista di stringhe / tuple (testo, opzioni)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    if isinstance(righe, str):
        righe = [righe]
    for i, r in enumerate(righe):
        opz = {}
        if isinstance(r, tuple):
            r, opz = r
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = opz.get("align", align)
        p.space_after = Pt(opz.get("spazio", spazio))
        puntato = opz.get("punto", False)
        run = p.add_run()
        run.text = ("•  " if puntato else "") + r
        f = run.font
        f.size = Pt(opz.get("size", size))
        f.bold = opz.get("bold", bold)
        f.color.rgb = opz.get("colore", colore)
        f.name = "Calibri"
    return tb


def punti(slide, x, y, w, h, voci, size=13, colore=TESTO):
    righe = []
    for v in voci:
        if isinstance(v, tuple):
            titolo, corpo = v
            righe.append((titolo, {"bold": True, "size": size, "colore": BLU, "spazio": 1}))
            righe.append((corpo, {"size": size - 1, "colore": colore, "spazio": 8}))
        else:
            righe.append((v, {"punto": True, "size": size, "spazio": 6}))
    return testo(slide, x, y, w, h, righe, size=size, colore=colore)


def immagine(slide, percorso, x, y, w=None, h=None, cornice=True):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    pic = slide.shapes.add_picture(percorso, Inches(x), Inches(y), **kw)
    if cornice:
        pic.line.color.rgb = RGBColor(0xD5, 0xDA, 0xE5)
        pic.line.width = Pt(0.75)
    return pic


def didascalia(slide, x, y, w, t):
    testo(slide, x, y, w, 0.3, t, size=10, colore=GRIGIO, align=PP_ALIGN.CENTER)


def nuova(numero, titolo, sottotitolo=None):
    s = prs.slides.add_slide(VUOTO)
    rett(s, 0, 0, 13.333, 0.95, NAVY)
    rett(s, 0, 0.95, 13.333, 0.06, ARANCIO)
    testo(s, 0.45, 0.14, 1.0, 0.7, f"{numero:02d}", size=28, colore=ARANCIO, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    testo(s, 1.35, 0.12, 9.5, 0.45, titolo, size=24, colore=BIANCO, bold=True)
    if sottotitolo:
        testo(s, 1.35, 0.55, 11.0, 0.35, sottotitolo, size=12, colore=RGBColor(0xC9, 0xD3, 0xF2))
    testo(s, 0.45, 7.1, 8.0, 0.3, f"Volley League · {STUDENTE} · {CORSO}", size=9, colore=GRIGIO)
    testo(s, 11.9, 7.1, 1.0, 0.3, f"{numero} / 10", size=9, colore=GRIGIO, align=PP_ALIGN.RIGHT)
    return s


def tabella(slide, x, y, w, righe, larghezze, size=11, h_riga=0.34):
    shape = slide.shapes.add_table(len(righe), len(righe[0]), Inches(x), Inches(y), Inches(w),
                                   Inches(h_riga * len(righe)))
    t = shape.table
    for j, lw in enumerate(larghezze):
        t.columns[j].width = Inches(lw)
    for i, r in enumerate(righe):
        t.rows[i].height = Inches(h_riga)
        for j, val in enumerate(r):
            c = t.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.06)
            c.margin_top = c.margin_bottom = Inches(0.02)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            c.fill.fore_color.rgb = NAVY if i == 0 else (CHIARO if i % 2 else BIANCO)
            p = c.text_frame.paragraphs[0]
            p.text = str(val)
            p.font.size = Pt(size)
            p.font.name = "Calibri"
            p.font.bold = i == 0
            p.font.color.rgb = BIANCO if i == 0 else TESTO
    return t


# ============================================================ 1 · Oggetto
s = prs.slides.add_slide(VUOTO)
rett(s, 0, 0, 13.333, 7.5, NAVY)
rett(s, 0, 5.3, 6.2, 0.06, ARANCIO)
testo(s, 0.6, 0.55, 6, 0.4, "01 · OGGETTO DEL PROGETTO", size=13, colore=ARANCIO, bold=True)
testo(s, 0.6, 1.0, 6.2, 1.6, ["Volley League", ("Piattaforma integrata per la gestione di campionati di pallavolo",
                                                {"size": 20, "bold": False, "colore": RGBColor(0xC9, 0xD3, 0xF2)})],
      size=44, colore=BIANCO, bold=True)
testo(s, 0.6, 2.75, 5.9, 2.5, [
    ("Sostituisce fogli Excel, referti cartacei e comunicati via email con un unico flusso digitale: "
     "dal tesseramento degli atleti all'omologazione del risultato, con classifica aggiornata in tempo reale.",
     {"size": 13, "colore": BIANCO, "spazio": 10}),
    ("4 ruoli: Giudice sportivo (admin) · Dirigente di società · Arbitro · Giocatore",
     {"size": 13, "colore": ARANCIO, "bold": True, "spazio": 6}),
    ("Dati demo: Serie D Maschile (10 squadre) e Serie C Maschile (8 squadre), girone A Emilia-Romagna",
     {"size": 12, "colore": BIANCO, "spazio": 6}),
    ("Angular 19 + Angular Material  ·  Spring Boot 3.4 + Spring Security JWT  ·  MySQL 8",
     {"size": 12, "colore": RGBColor(0xC9, 0xD3, 0xF2)}),
], size=13)
testo(s, 0.6, 5.5, 6, 1.4, [
    (STUDENTE, {"size": 18, "bold": True, "colore": BIANCO, "spazio": 2}),
    (MATRICOLA, {"size": 13, "colore": RGBColor(0xC9, 0xD3, 0xF2), "spazio": 2}),
    (CORSO, {"size": 13, "colore": RGBColor(0xC9, 0xD3, 0xF2), "spazio": 2}),
    ("Progetto individuale · 30 story point · 3 macro-funzionalità FE+BE+DB",
     {"size": 12, "colore": ARANCIO}),
])
immagine(s, img("01_login.png"), 6.95, 1.0, w=5.9)
immagine(s, img("03_classifica.png"), 7.85, 4.35, w=4.4)

# ============================================================ 2 · Team e ruoli
s = nuova(2, "Distribuzione del team e dei ruoli",
          "Progetto svolto in autonomia: tutti i ruoli Scrum ricoperti dallo stesso studente, con contratto da 30 SP")
ruoli = [("Product Owner", "Definisce visione, backlog e criteri di accettazione delle 3 macro-funzionalità"),
         ("Scrum Master", "Pianifica 3 sprint, conduce le cerimonie e aggiorna burndown e board Trello"),
         ("Sviluppatore DB", "Modello E-R, 10 tabelle MySQL, vincoli UNIQUE/FK, dump dello schema"),
         ("Sviluppatore BE", "API REST Spring Boot, sicurezza JWT/RBAC, pattern State/Strategy/Factory"),
         ("Sviluppatore FE", "SPA Angular Material responsive, guard per ruolo, validazione dei form"),
         ("Tester", "94 test JUnit 5 / Mockito / MockMvc, prove end-to-end sui 4 ruoli")]
for k, (r, d) in enumerate(ruoli):
    x = 0.45 + (k % 2) * 3.1
    y = 1.35 + (k // 2) * 1.75
    rett(s, x, y, 2.95, 1.6, CHIARO, MSO_SHAPE.ROUNDED_RECTANGLE)
    rett(s, x, y, 0.09, 1.6, ARANCIO)
    testo(s, x + 0.2, y + 0.12, 2.65, 0.4, r, size=15, colore=NAVY, bold=True)
    testo(s, x + 0.2, y + 0.55, 2.65, 1.0, d, size=11, colore=TESTO)
testo(s, 6.85, 1.3, 6.0, 0.4, "Contratto: 3 macro-funzionalità × (FE + BE + DB)", size=15, colore=NAVY, bold=True)
tabella(s, 6.85, 1.8, 6.0, [
    ["Macro-funzionalità", "FE", "BE", "DB", "SP"],
    ["MF1 · Organizzazione campionato e accesso", "2", "6", "2", "10"],
    ["MF2 · Gestione squadra e lista gara", "2", "4", "2", "8"],
    ["MF3 · Referto, omologazione e classifica", "3", "6", "3", "12"],
    ["Totale", "7", "16", "7", "30"],
], [3.6, 0.6, 0.6, 0.6, 0.6], size=12, h_riga=0.42)
testo(s, 6.85, 4.05, 6.0, 2.9, [
    ("Stima con sequenza di Fibonacci (1, 2, 3, 5, 8…)", {"bold": True, "colore": BLU, "spazio": 3}),
    ("Pivot: una CRUD completa = 2 SP. Task con logica di dominio non banale (validazione referto, "
     "calendario Berger, sicurezza JWT) stimati 3 SP.", {"size": 12, "spazio": 10}),
    ("Nessuna dipendenza esterna", {"bold": True, "colore": BLU, "spazio": 3}),
    ("Lavorando da solo, la Definition of Done di ogni card include codice commentato, test verdi e "
     "verifica manuale nel browser su desktop e mobile.", {"size": 12}),
], size=13)

# ============================================================ 3 · Cerimonie
s = nuova(3, "Evidenze delle cerimonie Scrum",
          "3 sprint da 5 giorni lavorativi · cerimonie svolte e annotate anche in forma individuale")
cer = [("Sprint planning", "Ad ogni inizio sprint: scelta delle card dal backlog Trello, stima in SP, "
                           "obiettivo di sprint scritto nella descrizione della lista."),
       ("Daily (log)", "Nota giornaliera di 3 righe su Trello: fatto ieri · farò oggi · impedimenti "
                      "(es. vincolo provincia arbitri, rollback referto)."),
       ("Sprint review", "Demo del flusso completato nel browser: S1 calendario + designazioni, "
                         "S2 rosa + lista gara, S3 referto → omologa → classifica."),
       ("Retrospettiva", "Start/Stop/Continue: es. 'scrivere prima i test del validatore set' (S2), "
                         "'tenere il seed demo allineato' (S3).")]
for k, (t, d) in enumerate(cer):
    y = 1.3 + k * 1.38
    rett(s, 0.45, y, 5.6, 1.25, CHIARO, MSO_SHAPE.ROUNDED_RECTANGLE)
    testo(s, 0.65, y + 0.1, 5.3, 0.4, t, size=15, colore=NAVY, bold=True)
    testo(s, 0.65, y + 0.48, 5.3, 0.8, d, size=11)
immagine(s, img("burndown.png"), 6.35, 1.3, w=6.55, cornice=False)
didascalia(s, 6.35, 4.95, 6.55, "Burndown dei 30 SP: le soste dopo ogni review coincidono con retro e planning")
tabella(s, 6.35, 5.35, 6.55, [
    ["Sprint", "Obiettivo", "SP"],
    ["S1", "MF1 completa: accesso, campionati, calendario, designazioni", "10"],
    ["S2", "MF2 completa + persistenza referto (T3.1)", "11"],
    ["S3", "Validazione referto, omologa, classifica, UI arbitro", "9"],
], [0.7, 5.15, 0.7], size=11, h_riga=0.36)

# ============================================================ 4 · Miro
s = nuova(4, "Disegno funzionale (Miro)",
          "Mappa delle 3 macro-funzionalità: ogni nota è una user story, il colore indica l'attore principale")
immagine(s, img("miro_disegno_funzionale.png"), 0.69, 1.25, w=11.95)

# ============================================================ 5 · Trello
s = nuova(5, "Board Trello", "12 card di sviluppo con etichetta di livello (FE/BE/DB) e story point · totale 30 SP")
immagine(s, img("trello_board.png"), 0.45, 1.25, w=9.6)
for k, (v, e) in enumerate([("12", "card di sviluppo"), ("30 SP", "stimati e completati"), ("3", "sprint da 5 giorni"),
                            ("100%", "card in Done")]):
    x = 0.45 + k * 2.43
    rett(s, x, 5.95, 2.3, 0.95, CHIARO, MSO_SHAPE.ROUNDED_RECTANGLE)
    testo(s, x + 0.15, 6.0, 0.95, 0.85, v, size=22, colore=BLU, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    testo(s, x + 1.1, 6.0, 1.15, 0.85, e, size=11, colore=TESTO, anchor=MSO_ANCHOR.MIDDLE)
punti(s, 10.25, 1.25, 2.7, 5.7, [
    ("Liste", "Backlog → Sprint 1/2/3 → Test & Documentazione"),
    ("Etichette", "FE (verde) · BE (viola) · DB (giallo) per verificare la copertura su tutti i livelli"),
    ("Story point", "Campo SP su ogni card; somma per macro-funzionalità 10 + 8 + 12"),
    ("Definition of Done", "Checklist 'test verdi' e 'verifica UI' spuntata prima di spostare in Done"),
    ("Import", "trello_backlog.csv in relazione_esame/ contiene l'intero backlog"),
], size=12)

# ============================================================ 6 · Test cases
s = nuova(6, "Test cases", "94 test automatici, 0 fallimenti · ./gradlew test (H2 in modalità MySQL, nessun DB esterno richiesto)")
tabella(s, 0.45, 1.3, 7.7, [
    ["Classe di test", "Tipo", "N.", "Cosa verifica"],
    ["RefertoValidatorTest", "Unit", "28", "Set a 25/15 con scarto 2, vincitore a 3 set"],
    ["ClassificaTest", "Unit", "14", "3-0/3-1 → 3-0 pt, 3-2 → 2-1, quozienti"],
    ["MacchinaStatiPartitaTest", "Unit", "10", "Transizioni lecite e illecite (State)"],
    ["GeneratoreCalendarioTest", "Unit", "11", "Berger: ogni coppia andata e ritorno"],
    ["VincoliEListaGaraTest", "Unit", "11", "Stessa società/provincia, 14 atleti, liberi"],
    ["DesignazioneServiceTest", "Mockito", "5", "Designazione rifiutata se vincolo violato"],
    ["RefertoTransazionaleIT", "Integr.", "6", "@Transactional rollback, firma, omologa"],
    ["SicurezzaRbacIT", "MockMvc", "8", "401 senza token, 403 per ruolo errato"],
    ["BackendApplicationTests", "Smoke", "1", "Avvio contesto Spring"],
], [2.35, 0.85, 0.45, 4.05], size=10.5, h_riga=0.4)
testo(s, 0.45, 5.45, 7.7, 1.5, [
    ("Esempio · rollback transazionale", {"bold": True, "colore": BLU, "spazio": 3}),
    ("Dato un referto con 3 set validi e una statistica riferita a un atleta non in lista gara, "
     "quando l'arbitro invia, allora il servizio lancia eccezione e nel DB non resta alcun set_partita: "
     "la partita torna IN_CORSO.", {"size": 11.5}),
], size=12)
immagine(s, img("30_test_report.png"), 8.4, 1.3, w=4.5)
didascalia(s, 8.4, 4.45, 4.5, "Report Gradle: 94 test, 100% successi")
testo(s, 8.4, 4.85, 4.5, 2.1, [
    ("Test manuali end-to-end", {"bold": True, "colore": BLU, "spazio": 3}),
    ("Percorso completo con i 4 account demo: designa → lista gara → referto con errore → correzione e "
     "firma → omologa → classifica e statistiche aggiornate. Verificato anche il layout mobile.",
     {"size": 11.5}),
])

# ============================================================ 7 · Architettura
s = nuova(7, "Disegno architetturale", "Architettura a 3 livelli, MVC lato backend, comunicazione REST/JSON stateless")
immagine(s, img("architettura.png"), 0.45, 1.25, h=5.75, cornice=False)
immagine(s, img("schema_er.png"), 7.75, 1.25, w=5.15, cornice=False)
didascalia(s, 7.75, 5.4, 5.15, "Schema E-R (database webapp, utente webappuser)")
testo(s, 7.75, 5.75, 5.15, 1.3, [
    ("Pattern: State (stato partita) · Strategy (criterio classifica: FIPAV, vittorie, quoziente set) · "
     "Factory (creazione set/statistiche e impronta SHA-256).", {"size": 11, "spazio": 4}),
    ("Errori uniformi via @RestControllerAdvice: 400 · 401 · 403 · 404 · 409 · 422.", {"size": 11}),
])

# ============================================================ 8 · Funzionalità implementate
s = nuova(8, "Funzionalità implementate", "Ogni ruolo vede solo il proprio menu; i controlli sono ripetuti lato server con @PreAuthorize")
colonne = [
    ("Giudice sportivo", "MF1 · MF3", ["Crea campionati e gironi", "Iscrive le squadre", "Genera calendario andata/ritorno (Berger)",
                                       "Designa arbitri con vincoli società, provincia, stesso giorno", "Omologa o respinge i referti"]),
    ("Dirigente", "MF2", ["Vede le squadre della propria società", "Tessera atleti e staff (CRUD rosa)",
                          "Compila lista gara: 6-14 atleti, max 2 liberi", "Invio liste → partita ROSTER_INVIATI"]),
    ("Arbitro", "MF3", ["Elenco designazioni", "Avvio gara", "Referto set per set con validazione live",
                        "Statistiche e cartellini per atleta", "Firma con password e invio"]),
    ("Giocatore", "MF3", ["Calendario della squadra", "Presenze e statistiche stagionali", "Grafico punti per gara",
                          "Classifica del girone"]),
]
for k, (ruolo, mf, voci) in enumerate(colonne):
    x = 0.45 + k * 3.13
    rett(s, x, 1.3, 2.98, 3.05, CHIARO, MSO_SHAPE.ROUNDED_RECTANGLE)
    rett(s, x, 1.3, 2.98, 0.5, BLU, MSO_SHAPE.ROUNDED_RECTANGLE)
    testo(s, x + 0.15, 1.33, 2.0, 0.45, ruolo, size=14, colore=BIANCO, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    testo(s, x + 1.95, 1.33, 0.95, 0.45, mf, size=10, colore=ARANCIO, bold=True, align=PP_ALIGN.RIGHT,
          anchor=MSO_ANCHOR.MIDDLE)
    testo(s, x + 0.12, 1.92, 2.78, 2.4, [(v, {"punto": True, "size": 11, "spazio": 4}) for v in voci])
for k, (f, d) in enumerate([("02_admin_campionati.png", "Campionati (giudice)"), ("06_designa_dialog.png", "Designazione con vincoli"),
                            ("09_rosa.png", "Rosa (dirigente)"), ("21_area_giocatore.png", "Area giocatore")]):
    x = 0.45 + k * 3.13
    immagine(s, img(f), x, 4.5, w=2.98)
    didascalia(s, x, 6.4, 2.98, d)
testo(s, 0.45, 6.7, 12.4, 0.35, "Accessibilità: lang=\"it\", label su ogni campo, contrasto AA, focus visibile, "
      "navigazione da tastiera, layout responsive con sidenav a scomparsa.", size=10.5, colore=GRIGIO)

# ============================================================ 9 · Dettaglio
s = nuova(9, "Dettaglio: referto digitale e omologazione (MF3)",
          "Il flusso più critico: deve garantire risultati regolari, atomicità del salvataggio e tracciabilità")
immagine(s, img("stati_partita.png"), 0.45, 1.2, w=7.0, cornice=False)
punti(s, 0.45, 2.95, 7.0, 4.1, [
    ("1 · Validazione (RefertoValidator)", "Set 1-4 a 25, tie-break a 15, sempre con 2 punti di scarto; chi vince "
                                           "deve arrivare esattamente a 3 set. Stesse regole replicate in Angular per l'errore immediato."),
    ("2 · Salvataggio atomico (@Transactional)", "Set, statistiche e cambio di stato in un'unica transazione: un atleta "
                                                 "non in lista gara provoca rollback completo (testato in RefertoTransazionaleIT)."),
    ("3 · Firma e impronta", "L'arbitro ri-digita la password (BCrypt); RefertoFactory calcola un hash SHA-256 del referto, "
                             "salvato in hash_referto per rilevare modifiche successive."),
    ("4 · Omologa → classifica", "Solo dopo l'omologa la Strategy FIPAV aggiorna la classifica: punti, poi quoziente set, poi quoziente punti."),
], size=12)
immagine(s, img("15_referto_errore_set.png"), 7.75, 1.2, w=5.15)
didascalia(s, 7.75, 4.45, 5.15, "Set irregolare (25-24) segnalato prima dell'invio")
immagine(s, img("17_referto_firma.png"), 7.75, 4.8, w=3.6)
testo(s, 11.45, 4.85, 1.5, 2.1, [("Firma digitale", {"size": 11, "bold": True, "colore": NAVY, "spazio": 4}),
                                 ("Conferma esplicita e password dell'arbitro prima dell'invio",
                                  {"size": 10, "colore": GRIGIO})])
rett(s, 0.45, 6.25, 7.0, 0.72, CHIARO, MSO_SHAPE.ROUNDED_RECTANGLE)
testo(s, 0.6, 6.29, 6.8, 0.65, [
    ("POST /api/partite/{id}/referto  ·  /omologa  ·  /respingi",
     {"size": 10.5, "bold": True, "colore": NAVY, "spazio": 2}),
    ("422 se i set non sono validi · 403 se l'arbitro non è quello designato · 409 se lo stato non lo consente",
     {"size": 10, "colore": GRIGIO})])

# ============================================================ 10 · Feedback
s = nuova(10, "Feedback", "Retrospettiva finale del progetto")
blocchi = [
    ("Cosa ha funzionato", VERDE, ["Stima in SP con pivot CRUD = 2 abbastanza precisa (scarto < 10%)",
                                   "Scrivere i test del validatore prima dell'UI ha evitato regressioni",
                                   "Il seed demo ha reso rapide review e verifica manuale"]),
    ("Difficoltà", RGBColor(0xB3, 0x26, 0x1E), ["Vincoli di designazione (provincia, stesso giorno) più complessi del previsto",
                                                 "Rollback con più entità: necessario flush esplicito e test di integrazione",
                                                 "Ricoprire tutti i ruoli Scrum da soli richiede disciplina nel log"]),
    ("Cosa ho imparato", BLU, ["Spring Security stateless con JWT e @PreAuthorize",
                               "Signals e componenti standalone in Angular 19",
                               "Applicare State, Strategy e Factory a regole di dominio reali"]),
    ("Sviluppi futuri", ARANCIO, ["Playoff e fasi finali", "Notifiche email delle designazioni",
                                  "Firma digitale qualificata e export PDF del referto"]),
]
for k, (t, c, voci) in enumerate(blocchi):
    x = 0.45 + (k % 2) * 6.25
    y = 1.3 + (k // 2) * 2.55
    rett(s, x, y, 6.1, 2.4, CHIARO, MSO_SHAPE.ROUNDED_RECTANGLE)
    rett(s, x, y, 0.1, 2.4, c)
    testo(s, x + 0.3, y + 0.15, 5.6, 0.45, t, size=17, colore=c, bold=True)
    testo(s, x + 0.3, y + 0.7, 5.6, 1.6, [(v, {"punto": True, "size": 13, "spazio": 6}) for v in voci])
testo(s, 0.45, 6.45, 12.4, 0.5, "Video demo: relazione_esame/demo.mp4  ·  Repository, Trello, Miro e Figma: relazione_esame/LINKS.txt",
      size=12, colore=NAVY, bold=True, align=PP_ALIGN.CENTER)

props = prs.core_properties
props.title = "Volley League - Relazione d'esame"
props.subject = CORSO
props.author = props.last_modified_by = STUDENTE
props.comments = ""
props.created = props.modified = datetime.datetime.now()
props.revision = 1

pptx_path = os.path.join(OUT, NOME + ".pptx")
prs.save(pptx_path)
print("creato", pptx_path, "-", len(prs.slides), "slide")

soffice = shutil.which("soffice") or shutil.which("libreoffice")
if soffice:
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", OUT, pptx_path], check=True,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    print("creato", os.path.join(OUT, NOME + ".pdf"))
